import csv
import io
import json
from calendar import month_name
from datetime import date

from fastapi import APIRouter, Depends, Query
from fastapi.responses import Response
from openpyxl import Workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app import models
from app.auth import get_current_user
from app.db import get_db
from app.services.queries import TxFilters, filtered_transactions, tx_filters

router = APIRouter(prefix="/api/export", tags=["export"])

HEADERS = ["Date", "Description", "Payer/Payee", "Amount", "Type", "Category", "Tags", "Account", "IBAN", "Notes"]


def _filename(f: TxFilters, ext: str) -> str:
    # Same naming as the original tracker: transactions_<year>_<Month>_<today>.csv
    name = "transactions"
    if f.year:
        name += f"_{f.year}"
    if f.month:
        name += f"_{month_name[f.month]}"
    return f"{name}_{date.today().isoformat()}.{ext}"


def _rows(db: Session, user: models.User, f: TxFilters):
    for t in db.scalars(filtered_transactions(user.id, f)).unique():
        yield t, [
            t.booking_date.isoformat(),
            t.description,
            t.payer,
            abs(t.amount_cents) / 100,
            "income" if t.amount_cents >= 0 else "expense",
            t.category.name if t.category else "",
            "; ".join(sorted(tag.name for tag in t.tags)),
            t.account.name,
            t.counterparty_iban,
            t.notes,
        ]


def _safe_cell(value):
    # Avoid spreadsheet formula injection from bank-provided text.
    if isinstance(value, str) and value[:1] in ("=", "+", "-", "@"):
        return "'" + value
    return value


@router.get("/csv")
def export_csv(f: TxFilters = Depends(tx_filters), db: Session = Depends(get_db), user: models.User = Depends(get_current_user)):
    buf = io.StringIO()
    writer = csv.writer(buf, delimiter=";", quoting=csv.QUOTE_MINIMAL)
    writer.writerow(HEADERS)
    for _, row in _rows(db, user, f):
        row[3] = f"{row[3]:.2f}".replace(".", ",")  # European decimal comma, like the original export
        writer.writerow([_safe_cell(c) for c in row])
    return Response(
        "﻿" + buf.getvalue(),
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": f'attachment; filename="{_filename(f, "csv")}"'},
    )


@router.get("/xlsx")
def export_xlsx(f: TxFilters = Depends(tx_filters), db: Session = Depends(get_db), user: models.User = Depends(get_current_user)):
    wb = Workbook()
    ws = wb.active
    ws.title = "Transactions"
    ws.append(HEADERS)
    for t, row in _rows(db, user, f):
        row[0] = t.booking_date
        row[3] = t.amount_cents / 100  # signed in Excel so SUM() works
        ws.append([_safe_cell(c) for c in row])
    for cell in ws["A"][1:]:
        cell.number_format = "DD.MM.YYYY"
    for cell in ws["D"][1:]:
        cell.number_format = "#,##0.00 €"
    ws.freeze_panes = "A2"
    for col, width in zip("ABCDEFGHIJ", (12, 50, 30, 12, 9, 18, 20, 16, 26, 30)):
        ws.column_dimensions[col].width = width
    buf = io.BytesIO()
    wb.save(buf)
    return Response(
        buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{_filename(f, "xlsx")}"'},
    )


@router.get("/json")
def export_json(db: Session = Depends(get_db), user: models.User = Depends(get_current_user)):
    """Full backup of the current user's data."""

    def dump(model, **extra):
        cols = [c.name for c in model.__table__.columns if c.name not in ("user_id", "password_hash")]
        out = []
        for obj in db.scalars(select(model).where(model.user_id == user.id)):
            row = {c: getattr(obj, c) for c in cols}
            row.update({k: fn(obj) for k, fn in extra.items()})
            out.append(row)
        return out

    data = {
        "version": 1,
        "exported_at": date.today().isoformat(),
        "username": user.username,
        "accounts": dump(models.Account),
        "categories": dump(models.Category),
        "transactions": dump(models.Transaction, tags=lambda t: sorted(tag.name for tag in t.tags)),
        "rules": dump(models.Rule),
        "budgets": dump(models.Budget),
        "recurring": dump(models.RecurringSeries),
    }
    body =json.dumps(data, default=str, ensure_ascii=False, indent=1)
    return Response(
        body,
        media_type="application/json",
        headers={"Content-Disposition": f'attachment; filename="financetracker_backup_{date.today().isoformat()}.json"'},
    )
